"""Convertit le classeur d'analyse (onglets Documents + Coordonnées) en annotations JSON.

Usage : python -m eval.xlsx_to_annotations CLASSEUR.xlsx data/annotations
Les annotations produites sont marquées verified=false : elles viennent d'une lecture visuelle
faite par Claude et doivent être relues avant d'être utilisées comme référence.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from model.annotation import AnnotatedPoint, Annotation, save_annotation
from model.schema import Area

ABSENT = {"", "(vide)", "non visible", "non lisible", "non mentionné"}


def _txt(v: object) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return None if s.lower() in ABSENT else s


def convert(xlsx: Path, out_dir: Path, annotator: str = "claude-lecture-visuelle") -> list[Path]:
    if not xlsx.is_file() or xlsx.suffix.lower() != ".xlsx":
        raise SystemExit(f"Il faut le chemin d'un fichier .xlsx, pas : {xlsx}")
    wb = load_workbook(xlsx, data_only=True)
    docs = {}
    for row in wb["Documents"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            docs[row[0]] = row
    decimals = {k: row[11] for k, row in docs.items()}

    points: dict[str, list[AnnotatedPoint]] = {}
    for row in wb["Coordonnées"].iter_rows(min_row=2, values_only=True):
        did, label, x, y = row[0], row[1], row[2], row[3]
        if did is None or x is None or y is None:
            continue
        d = decimals.get(did)
        if d is None:
            raise ValueError(f"{did}: nombre de décimales inconnu, impossible de formater {label}")
        points.setdefault(did, []).append(AnnotatedPoint(label=str(label), x=f"{x:.{d}f}", y=f"{y:.{d}f}"))

    written = []
    for did, row in docs.items():
        if did not in points:   # document non testé : pas de tableau exploitable, pas d'annotation
            continue
        a = Annotation(
            document_id=did,
            image=str(row[1]),
            locality=_txt(row[3]),
            title_number=_txt(row[4]),
            nicad=_txt(row[5]),
            printed_area=Area(ha=int(row[6] or 0), a=int(row[7] or 0), ca=int(row[8] or 0)),
            decimals=decimals[did],
            quality=_txt(row[15]),
            points=points[did],
            annotator=annotator,
            notes=_txt(row[17]),
        )
        written.append(save_annotation(a, out_dir))
    return written


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    for p in convert(Path(sys.argv[1]), Path(sys.argv[2])):
        print("écrit", p)
